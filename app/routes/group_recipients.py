import logging
import csv
import io
from io import BytesIO
from datetime import datetime, timezone
from fastapi import APIRouter, Depends, Request, Form, Path, File, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from typing import Optional
from sqlalchemy.exc import IntegrityError
from sqlalchemy import func, or_

from app.db.database import get_db
from app.models.recipient import GroupRecipient
from app.models.camera_group import CameraGroup
from app.models.camera import Camera
from app.models.camera_email_notification_log import CameraEmailNotificationLog, CameraEmailNotificationRecipient
from app.utils.audit_logger import log_audit
from app.utils.template_helper import templates
from app.utils.email_helper import send_email, build_email_body
from app.utils.response_helper import json_error_response, json_success_response
from app.utils.smtp_config import get_smtp_status_message
from app.routes.auth import admin_access_required
from app.models.user import User

logger = logging.getLogger("main")
mgmt_logger = logging.getLogger("management")

router = APIRouter()


# ============================================================
# 1️⃣ RENDER HALAMAN HTML
# ============================================================
@router.get("/admin/recipients", name="admin_list_recipients_page")
@router.get("/recipients", name="list_recipients_page")
def list_recipients_page(request: Request, db: Session = Depends(get_db)):
    groups = db.query(CameraGroup).all()
    smtp_status = get_smtp_status_message()
    return templates.TemplateResponse("recipients.html", {
        "request": request, 
        "groups": groups,
        "smtp_status": smtp_status
    })


# ============================================================
# 2️⃣ API: GET RECIPIENTS (JSON) - FIXED WITH SEARCH & FILTERS
# ============================================================
@router.get("/api/recipients", name="list_recipients_api")
def list_recipients_api(
    db: Session = Depends(get_db), 
    page: int = 1, 
    per_page: int = 10,
    search: Optional[str] = None,
    group: Optional[str] = None,
    location: Optional[str] = None
):
    try:
        groups = db.query(CameraGroup).all()
        
        # Base query with join
        query = db.query(GroupRecipient).options(joinedload(GroupRecipient.group))
        
        # Apply search filter (email or nickname)
        if search:
            search_filter = f"%{search}%"
            query = query.filter(
                or_(
                    GroupRecipient.email.ilike(search_filter),
                    GroupRecipient.nickname.ilike(search_filter)
                )
            )
        
        # Apply group filter
        if group and group.isdigit():
            query = query.filter(GroupRecipient.group_id == int(group))
        
        # Apply location filter
        if location:
            location_filter = f"%{location}%"
            query = query.filter(
                or_(
                    GroupRecipient.locations.ilike(location_filter),
                    GroupRecipient.locations.is_(None) if location == "" else False
                )
            )
        
        # Get total count after filters
        total = query.count()
        
        # Apply pagination
        recipients = query.offset((page - 1) * per_page).limit(per_page).all()
        total_pages = (total + per_page - 1) // per_page

        # lokasi unik per group
        locations_by_group = {}
        for g in groups:
            locs = (
                db.query(Camera.location)
                .filter(Camera.groups.any(CameraGroup.id == g.id))
                .distinct()
                .all()
            )
            locations_by_group[g.id] = [l[0] for l in locs if l[0]]

        return json_success_response(
            "Recipient data fetched successfully",
            {
                "recipients": [
                    {
                        "id": r.id,
                        "email": r.email,
                        "nickname": r.nickname,
                        "group_id": r.group_id,
                        "group_name": r.group.name if r.group else None,
                        "locations": r.locations,
                    }
                    for r in recipients
                ],
                "groups": [{"id": g.id, "name": g.name} for g in groups],
                "locations_by_group": locations_by_group,
                "pagination": {
                    "page": page,
                    "per_page": per_page,
                    "total": total,
                    "total_pages": total_pages,
                },
            },
        )

    except Exception as e:
        logger.exception("Error fetching recipients JSON")
        return json_error_response(str(e), 500)


# ============================================================
# 3️⃣ ADD RECIPIENT
# ============================================================
@router.post("/recipients/add", name="add_recipient")
async def add_recipient(
    request: Request,
    email: str = Form(...),
    nickname: Optional[str] = Form(None),
    group_id: int = Form(...),
    db: Session = Depends(get_db),
):
    form = await request.form()
    locs = form.getlist("locations")
    locations_str = ",".join(locs) if locs else None

    try:
        recipient = GroupRecipient(
            email=email.strip(),
            nickname=nickname,
            group_id=group_id,
            locations=locations_str,
        )
        db.add(recipient)
        db.flush()
        db.commit()
        db.refresh(recipient)
        return json_success_response("Recipient added successfully", {"id": recipient.id})
    except IntegrityError:
        db.rollback()
        logger.exception("IntegrityError on add_recipient")
        return json_error_response(f"Email '{email}' already exist in this group", 400)
    except Exception as e:
        db.rollback()
        logger.exception("Unexpected error on add_recipient")
        return json_error_response(str(e), 500)


# ============================================================
# 4️⃣ EDIT RECIPIENT
# ============================================================
@router.post("/recipients/edit/{recipient_id}")
async def edit_recipient_submit(
    request: Request,
    recipient_id: int = Path(...),
    db: Session = Depends(get_db),
    email: str = Form(...),
    nickname: str = Form(None),
    group_id: int = Form(...),
):
    rec = db.query(GroupRecipient).filter(GroupRecipient.id == recipient_id).first()
    if not rec:
        return json_error_response("Recipient not found", 404)

    form = await request.form()
    locs = form.getlist("locations")
    locations_str = ",".join(locs) if locs else None
    email_norm = email.strip()

    before = {
        "email": rec.email,
        "nickname": rec.nickname,
        "group_id": rec.group_id,
        "locations": rec.locations,
    }

    try:
        rec.email = email_norm
        rec.nickname = nickname
        rec.group_id = group_id
        rec.locations = locations_str
        db.flush()
        db.commit()
        db.refresh(rec)
    except IntegrityError:
        db.rollback()
        logger.exception("IntegrityError on edit_recipient")
        return json_error_response(f"Email '{email_norm}' already exist in this group", 400)
    except Exception as e:
        db.rollback()
        logger.exception("Unexpected error on edit_recipient")
        return json_error_response(str(e), 500)

    after = {
        "email": rec.email,
        "nickname": rec.nickname,
        "group_id": rec.group_id,
        "locations": rec.locations,
    }
    changes = [f"- {k}: '{before[k]}' -> '{after[k]}'" for k in before if before[k] != after[k]]

    log_audit(
        db=db,
        user=request.session.get("user_name", "unknown"),
        action="update_recipient",
        target=rec.email,
        ip=request.client.host,
        extra="\n".join(changes) if changes else "No changes detected.",
    )

    return json_success_response("Recipient updated successfully")


# ============================================================
# 5️⃣ DELETE RECIPIENT
# ============================================================
@router.post("/recipients/delete/{recipient_id}")
async def delete_recipient_submit(
    request: Request,
    recipient_id: int = Path(...),
    db: Session = Depends(get_db),
):
    try:
        rec = db.query(GroupRecipient).filter(GroupRecipient.id == recipient_id).first()
        if not rec:
            return json_error_response("Recipient not found", 404)

        target_email = rec.email
        db.delete(rec)
        db.commit()

        log_audit(
            db=db,
            user=request.session.get("user_name", "unknown"),
            action="delete_recipient",
            target=target_email,
            ip=request.client.host,
            extra="Recipient deleted",
        )

        return json_success_response(f"Recipient {target_email} deleted")
    except Exception as e:
        db.rollback()
        logger.exception("Unexpected error on delete_recipient")
        return json_error_response(str(e), 500)


# ============================================================
# 6️⃣ TEST SEND EMAIL (GROUP + LOCATION) — LOG KE DB
# ============================================================
@router.post("/api/recipients/test-send")
def test_send_email(
    group_id: int = Form(...),
    location: Optional[str] = Form(None),
    db: Session = Depends(get_db),
):
    """
    Kirim test email ke semua recipient berdasarkan group_id dan lokasi (opsional).
    Log dicatat ke CameraEmailNotificationLog agar konsisten dengan alert sungguhan.
    ✅ UPDATE: CC email juga masuk ke log
    """

    try:
        group = db.query(CameraGroup).filter(CameraGroup.id == group_id).first()
        if not group:
            return json_error_response(f"Group ID {group_id} not found.", 404)

        # Ambil recipient berdasarkan lokasi (opsional)
        q = db.query(GroupRecipient).filter(GroupRecipient.group_id == group_id)
        if location:
            q = q.filter(func.lower(GroupRecipient.locations).ilike(f"%{location.lower()}%"))
        recipients = q.all()
        emails = [r.email for r in recipients if r.email]

        if not emails:
            logger.warning(
                "No recipients found for test email in group '%s' (location: %s)",
                group.name, location or "ALL",
            )
            return json_error_response(
                f"No recipients for group'{group.name}' location '{location or 'ALL'}'", 404
            )

        group_name = group.name
        subject = f"🔔 B-SNAP Test Email — {group_name}{f' | {location}' if location else ''}"
        now_utc = datetime.now(timezone.utc)
        now_str = now_utc.strftime("%d %b %Y %H:%M:%S")

        plain_body, html_body = build_email_body(
            camera_name=f"TEST — {group_name}",
            ip="127.0.0.1",
            asset_no="N/A",
            coordinate="N/A",
            incident_time=now_str,
            last_snapshot_time=now_str,
            has_snapshot=False,
        )

        # --- Buat log baru sebelum kirim ---
        log_entry = CameraEmailNotificationLog(
            camera_id=None,  # dummy camera_id
            camera_name=f"TEST_{group_name}",
            incident_started_at=now_utc,
            sent_at=now_utc,
            error_message=None,
            success=False,
            type="test",
        )
        db.add(log_entry)
        db.flush()  # agar dapat log_entry.id

        # ✅ ROBUST LOGGING: Immutable recipient snapshot (To + CC)
        for em in emails:
            db.add(CameraEmailNotificationRecipient(
                log_id=log_entry.id,
                recipient_email=em,
            ))
        
        # ✅ TAMBAHAN: Log CC recipient dari config
        from app.utils.smtp_config import get_smtp_config
        config = get_smtp_config()
        cc_email = config.get("email_cc")
        if cc_email and isinstance(cc_email, str) and cc_email.strip():
            db.add(CameraEmailNotificationRecipient(
                log_id=log_entry.id,
                recipient_email=cc_email.strip()
            ))
            logger.info(f"Added CC recipient to test email log: {cc_email.strip()}")
        
        db.flush()
        db.commit()  # commit sementara agar log aman disimpan

        # --- Kirim email ---
        try:
            send_email(emails, subject, plain_body, html_body)

            log_entry.success = True
            log_entry.error_message = None
            log_entry.sent_at = datetime.now(timezone.utc)
            db.commit()

            logger.info(
                "✅ Test email sent successfully for group '%s' (%s) to %s",
                group_name, location or "ALL", emails,
            )
            return json_success_response(
                f"Test email sent successfully to {len(emails)} recipient(s) in group '{group_name}'",
                {"recipients": emails, "log_id": log_entry.id},
            )

        except Exception as e:
            db.rollback()
            log_entry.success = False
            log_entry.error_message = str(e)
            db.commit()

            logger.exception(
                "❌ Failed to send test email for group '%s' (location: %s): %s",
                group_name, location or "ALL", e,
            )
            return json_error_response(f"Failed sent email: {e}", 500)

    except Exception as e:
        db.rollback()
        logger.exception("Unexpected error on test_send_email")
        return json_error_response(str(e), 500)


# ============================================================
# 7️⃣ EXPORT RECIPIENTS TO CSV
# ============================================================
@router.get("/api/recipients/export", name="export_recipients_csv")
def export_recipients_csv(
    db: Session = Depends(get_db),
    current_admin: User = Depends(admin_access_required)
):
    """Export all recipients to CSV."""
    try:
        recipients = db.query(GroupRecipient).options(joinedload(GroupRecipient.group)).all()
        
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["email", "nickname", "group_name", "locations"])
        
        for r in recipients:
            writer.writerow([
                r.email,
                r.nickname or "",
                r.group.name if r.group else "",
                r.locations or ""
            ])
        
        csv_content = output.getvalue()
        output.close()
        
        return StreamingResponse(
            io.BytesIO(csv_content.encode('utf-8')),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=recipients_export.csv"}
        )
    except Exception as e:
        logger.exception("Error exporting recipients")
        return json_error_response(str(e), 500)


# ============================================================
# 8️⃣ BULK IMPORT RECIPIENTS FROM CSV
# ============================================================
@router.post("/api/recipients/import", name="import_recipients_csv")
async def import_recipients_csv(
    request: Request,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_admin: User = Depends(admin_access_required)
):
    """Bulk import recipients from CSV file.
    
    CSV Format: email,nickname,group_name,locations
    - locations can be comma-separated within the column
    """
    try:
        if not file.filename.endswith('.csv'):
            return json_error_response("File must be a CSV", 400)
        
        content = await file.read()
        content_str = content.decode('utf-8')
        
        csv_file = io.StringIO(content_str)
        reader = csv.DictReader(csv_file)
        
        required_columns = {'email', 'group_name'}
        if not required_columns.issubset(reader.fieldnames or []):
            return json_error_response(f"CSV must have columns: {required_columns}", 400)
        
        # Cache groups for lookup
        groups = {g.name: g.id for g in db.query(CameraGroup).all()}
        
        imported = 0
        skipped = 0
        errors = []
        
        for row_num, row in enumerate(reader, start=2):  # start=2 for human-friendly line number
            try:
                email = row.get('email', '').strip()
                nickname = row.get('nickname', '').strip() or None
                group_name = row.get('group_name', '').strip()
                locations = row.get('locations', '').strip() or None
                
                if not email or not group_name:
                    errors.append(f"Row {row_num}: Missing email or group_name")
                    skipped += 1
                    continue
                
                # Find group_id
                group_id = groups.get(group_name)
                if not group_id:
                    errors.append(f"Row {row_num}: Group '{group_name}' not found")
                    skipped += 1
                    continue
                
                # Check if recipient already exists
                existing = db.query(GroupRecipient).filter(
                    GroupRecipient.email == email,
                    GroupRecipient.group_id == group_id
                ).first()
                
                if existing:
                    skipped += 1
                    continue
                
                # Create new recipient
                recipient = GroupRecipient(
                    email=email,
                    nickname=nickname,
                    group_id=group_id,
                    locations=locations
                )
                db.add(recipient)
                imported += 1
                
            except Exception as e:
                errors.append(f"Row {row_num}: {str(e)}")
                skipped += 1
        
        db.commit()
        
        # Log audit
        log_audit(
            db=db,
            user=request.session.get("user_name", "unknown"),
            action="bulk_import_recipients",
            target="recipients",
            ip=request.client.host,
            extra=f"Imported: {imported}, Skipped: {skipped}, Errors: {len(errors)}"
        )
        
        return json_success_response(
            f"Import complete: {imported} imported, {skipped} skipped",
            {"imported": imported, "skipped": skipped, "errors": errors[:10]}  # Limit errors in response
        )
        
    except Exception as e:
        db.rollback()
        logger.exception("Error importing recipients")
        return json_error_response(str(e), 500)


# ============================================================
# 9️⃣ GET LOCATIONS WITHOUT RECIPIENTS (JSON)
# ============================================================
@router.get("/api/recipients/locations-without-recipients")
def get_locations_without_recipients(
    db: Session = Depends(get_db),
    group_id: Optional[int] = None
):
    """
    Get all locations yang tidak memiliki email recipient untuk group tertentu.
    
    Jika group_id tidak diberikan, return semua locations di seluruh system.
    """
    try:
        if group_id:
            # Get all locations untuk group tertentu
            all_locations = db.query(Camera.location)\
                .filter(Camera.groups.any(CameraGroup.id == group_id), Camera.location.isnot(None))\
                .distinct()\
                .all()
            
            # Get locations yang sudah memiliki recipient
            covered_locations = db.query(func.lower(GroupRecipient.locations))\
                .filter(GroupRecipient.group_id == group_id, GroupRecipient.locations.isnot(None))\
                .all()
        else:
            # Get semua locations di seluruh system
            all_locations = db.query(Camera.location)\
                .filter(Camera.location.isnot(None))\
                .distinct()\
                .all()
            
            # Get semua locations dari recipients
            covered_locations = db.query(func.lower(GroupRecipient.locations))\
                .filter(GroupRecipient.locations.isnot(None))\
                .all()
        
        # Parse locations dari recipient (comma-separated)
        covered_set = set()
        for loc_tuple in covered_locations:
            if loc_tuple[0]:
                # Split comma-separated locations dan normalize
                locations = [l.strip() for l in loc_tuple[0].split(',')]
                covered_set.update(l.lower() for l in locations)
        
        # Find locations without recipients
        uncovered_locations = []
        for loc_tuple in all_locations:
            if loc_tuple[0]:  # Skip NULL locations
                loc_normalized = loc_tuple[0].lower().strip()
                if loc_normalized not in covered_set:
                    uncovered_locations.append(loc_tuple[0])
        
        # Count cameras per uncovered location
        location_camera_count = {}
        for loc in uncovered_locations:
            if group_id:
                count = db.query(func.count(Camera.id))\
                    .filter(Camera.groups.any(CameraGroup.id == group_id), Camera.location == loc)\
                    .scalar()
            else:
                count = db.query(func.count(Camera.id))\
                    .filter(Camera.location == loc)\
                    .scalar()
            location_camera_count[loc] = count
        
        # Sort by location name
        uncovered_locations = sorted(uncovered_locations)
        
        return json_success_response(
            "Locations without recipients fetched successfully",
            {
                "locations": uncovered_locations,
                "location_camera_count": location_camera_count,
                "total_uncovered": len(uncovered_locations),
            }
        )
    
    except Exception as e:
        logger.exception("Error fetching locations without recipients")
        return json_error_response(str(e), 500)


# ============================================================
# 🔟 EXPORT LOCATIONS WITHOUT RECIPIENTS TO EXCEL
# ============================================================
@router.get("/api/recipients/export-locations-without-recipients")
def export_locations_without_recipients_excel(
    request: Request,
    db: Session = Depends(get_db),
    group_id: Optional[int] = None,
    current_admin: User = Depends(admin_access_required)
):
    """Export locations without recipients to Excel file."""
    try:
        # Import openpyxl untuk Excel
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        
        if group_id:
            # Get group name
            group = db.query(CameraGroup).filter(CameraGroup.id == group_id).first()
            group_name = group.name if group else "Unknown"
            
            # Get all locations untuk group tertentu
            all_locations = db.query(Camera.location)\
                .filter(Camera.groups.any(CameraGroup.id == group_id), Camera.location.isnot(None))\
                .distinct()\
                .all()
            
            # Get locations yang sudah memiliki recipient
            covered_locations = db.query(func.lower(GroupRecipient.locations))\
                .filter(GroupRecipient.group_id == group_id, GroupRecipient.locations.isnot(None))\
                .all()
        else:
            group_name = "All Groups"
            # Get semua locations di seluruh system
            all_locations = db.query(Camera.location)\
                .filter(Camera.location.isnot(None))\
                .distinct()\
                .all()
            
            # Get semua locations dari recipients
            covered_locations = db.query(func.lower(GroupRecipient.locations))\
                .filter(GroupRecipient.locations.isnot(None))\
                .all()
        
        # Parse locations dari recipient
        covered_set = set()
        for loc_tuple in covered_locations:
            if loc_tuple[0]:
                locations = [l.strip() for l in loc_tuple[0].split(',')]
                covered_set.update(l.lower() for l in locations)
        
        # Find locations without recipients
        uncovered_locations = []
        for loc_tuple in all_locations:
            if loc_tuple[0]:
                loc_normalized = loc_tuple[0].lower().strip()
                if loc_normalized not in covered_set:
                    uncovered_locations.append(loc_tuple[0])
        
        # Count cameras per location
        location_camera_count = {}
        for loc in uncovered_locations:
            if group_id:
                count = db.query(func.count(Camera.id))\
                    .filter(Camera.groups.any(CameraGroup.id == group_id), Camera.location == loc)\
                    .scalar()
            else:
                count = db.query(func.count(Camera.id))\
                    .filter(Camera.location == loc)\
                    .scalar()
            location_camera_count[loc] = count
        
        # Create Excel workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Locations Without Recipients"
        
        # Define styles
        header_fill = PatternFill(start_color="FF3B82F6", end_color="FF3B82F6", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        
        center_alignment = Alignment(horizontal="center", vertical="center")
        left_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        
        # Add title and metadata
        ws.merge_cells("A1:C1")
        title_cell = ws["A1"]
        title_cell.value = f"📍 Locations Without Email Recipients - {group_name}"
        title_cell.font = Font(bold=True, size=14, color="FF1F2937")
        title_cell.alignment = center_alignment
        
        ws.merge_cells("A2:C2")
        timestamp_cell = ws["A2"]
        timestamp_cell.value = f"Generated: {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M:%S UTC')}"
        timestamp_cell.font = Font(size=10, color="FF6B7280", italic=True)
        timestamp_cell.alignment = center_alignment
        
        # Add headers
        ws.append([])  # Empty row
        headers = ["#", "Location", "Camera Count"]
        ws.append(headers)
        
        # Style headers
        for cell in ws[4]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = border
        
        # Add data
        sorted_locations = sorted(uncovered_locations)
        for idx, location in enumerate(sorted_locations, 1):
            camera_count = location_camera_count.get(location, 0)
            ws.append([idx, location, camera_count])
            
            # Style data rows
            for cell in ws[4 + idx]:
                cell.border = border
                if cell.column == 1:  # Index column
                    cell.alignment = center_alignment
                elif cell.column == 3:  # Camera count column
                    cell.alignment = center_alignment
                else:  # Location column
                    cell.alignment = left_alignment
        
        # Set column widths
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 40
        ws.column_dimensions["C"].width = 15
        
        # Add summary section
        summary_row = 4 + len(sorted_locations) + 2
        ws[f"A{summary_row}"] = "Total Locations Without Recipients:"
        ws[f"A{summary_row}"].font = Font(bold=True)
        ws[f"B{summary_row}"] = len(sorted_locations)
        ws[f"B{summary_row}"].font = Font(bold=True)
        
        ws[f"A{summary_row + 1}"] = "Total Cameras in These Locations:"
        ws[f"A{summary_row + 1}"].font = Font(bold=True)
        ws[f"B{summary_row + 1}"] = sum(location_camera_count.values())
        ws[f"B{summary_row + 1}"].font = Font(bold=True)
        
        # Prepare response
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Log audit
        log_audit(
            db=db,
            user=request.session.get("user_name", "unknown") if hasattr(request, 'session') else "api_user",
            action="export_locations_without_recipients",
            target=group_name,
            ip=request.client.host if hasattr(request, 'client') else "N/A",
            extra=f"Exported {len(sorted_locations)} locations without recipients"
        )
        
        filename = f"locations_without_recipients_{group_name.replace(' ', '_')}_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    
    except Exception as e:
        logger.exception("Error exporting locations without recipients to Excel")
        return json_error_response(str(e), 500)
